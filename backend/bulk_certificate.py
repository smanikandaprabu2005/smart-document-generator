import os
import csv
import io
import tempfile
import openpyxl
import re
import uuid
import zipfile
import logging
import multiprocessing
from concurrent.futures import ThreadPoolExecutor
from flask import request, send_file, jsonify
from backend.utils.pdf_generator import create_document
from backend.db_utils import store_certificate
from datetime import datetime

logger = logging.getLogger(__name__)

def extract_year_and_branch(name):
    match = re.search(r'\((\d{2})([A-Z]{3})\d+\)', name)
    if not match:
        return "", ""

    adm_year = int(match.group(1))      # e.g., 21
    branch_code = match.group(2)        # e.g., BAD

    # ------- GET CURRENT YEAR DYNAMICALLY -------
    # Example: 2025 → 25
    current_year_full = datetime.now().year      # 2025
    current_year = int(str(current_year_full)[-2:])   # 25

    # ------- YEAR CALCULATION -------
    year_no = current_year - adm_year + 1        # (25 - 21 + 1 = 5)

    if year_no > 4:
        year = "passed out"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(year_no, "th")
        year = f"{year_no}{suffix} year"

    # ------- BRANCH MAPPING -------
    branch_map = {
        "BAD": "AIDS",
        "BCS": "CSE",
        "BEC": "ECE",
        "BME": "MECH",
        "BEE": "EEE",
        "BIT": "IT"
    }
    branch = branch_map.get(branch_code, "")

    return year, branch


def generate_single_certificate(row, name):
    """Worker function to generate a single certificate.

    Runs in parallel thread pool.
    """
    doc_id = str(uuid.uuid4())
    year, branch = extract_year_and_branch(name)

    logger.info(f"Generating certificate for {name} (doc_id={doc_id})")

    placeholders = {
        'name': name,
        'event': row.get('Event') or row.get('Event Name', ''),
        'date': row.get('date') or row.get('From Date', ''),
        'role': row.get('Role') or 'presented',
        'organizer': row.get('Club Name', ''),
        'year': year,
        'branch': branch,
    }

    template_path = os.path.join(
        os.path.dirname(__file__),
        'default_templates',
        'certificate_template.docx'
    )

    pdf_buffer = create_document(
        doc_type='certificate',
        template_path=template_path,
        placeholders=placeholders,
        doc_id=doc_id
    )

    # Save PDF to temp file
    temp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
    temp.write(pdf_buffer.read())
    temp.close()

    # Store certificate metadata in MongoDB
    store_certificate(doc_id, name, row.get('Event', ''), row.get('Date', ''), row.get('Role', ''), 'certificate')

    return temp.name


def process_bulk_certificates():
    """
    Endpoint to handle bulk certificate generation from CSV or Excel file.
    Each row may have multiple participant names (comma or semicolon separated).
    """
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    file = request.files['file']
    filename = file.filename
    if filename.endswith('.csv'):
        content = file.read().decode('utf-8')
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)
        logger.info(f"Parsed {len(rows)} rows from CSV")
    elif filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h) for h in rows[0]]
        rows = [dict(zip(headers, row)) for row in rows[1:]]
        logger.info(f"Parsed {len(rows)} rows from Excel")
    else:
        return jsonify({'error': 'Unsupported file type'}), 400

    generated_files = []
    # Collect all certificate generation tasks
    tasks = []
    for row in rows:
        names_field = row.get('Names') or row.get('Name') or row.get('Recipient Name') or row.get('Student Coordinators/Presenters')
        logger.debug(f"Row data: {row}")
        logger.debug(f"names_field: {names_field}")
        if not names_field:
            continue
        names = [n.strip() for n in re.split(r'[;,]', names_field) if n.strip()]
        logger.debug(f"Parsed names: {names}")
        # Add all (row, name) pairs to task list
        for name in names:
            tasks.append((row, name))
    
    # Generate certificates in parallel using ThreadPoolExecutor
    max_workers = min(4, multiprocessing.cpu_count())
    logger.info(f"Using ThreadPoolExecutor with {max_workers} workers")

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        for i, file_path in enumerate(executor.map(lambda x: generate_single_certificate(*x), tasks)):
            logger.info(f"Processed certificate {i + 1}/{len(tasks)}: {file_path}")
            generated_files.append(file_path)
    # Zip all PDFs
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w') as zipf:
        for file_path in generated_files:
            zipf.write(file_path, os.path.basename(file_path))
    # Clean up temp files
    for file_path in generated_files:
        os.remove(file_path)
    zip_buffer.seek(0)
    return send_file(zip_buffer, as_attachment=True, download_name='certificates.zip', mimetype='application/zip')
